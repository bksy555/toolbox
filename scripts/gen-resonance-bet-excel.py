#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 共振期投冷号 策略 300 期投注明细 Excel"""
import datetime, json, os, io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GAN_TO_DAN = {"甲":[1,4,8],"乙":[3,4,8],"丙":[3,4,9],"丁":[2,4,9],"戊":[3,4,9],"己":[2,4,9],"庚":[2,7,9],"辛":[2,6,7],"壬":[1,6,7],"癸":[1,6,8]}
TIAN_GAN = ["甲","乙","丙","丁","戊","己","庚","辛","壬","癸"]
WEEK_CN = ["", "周一", "周二", "周三", "周四", "周五", "周六", "周日"]

def get_hai_hour_gan(y,m,d):
    diff = (datetime.date(y,m,d)-datetime.date(1900,1,1)).days
    cycle = ((diff+10)%60+60)%60
    return TIAN_GAN[(((cycle%10)%5)*2+11)%10]
def dans_for(y,m,d): return GAN_TO_DAN[get_hai_hour_gan(y,m,d)]
def period_to_date(p):
    return datetime.date(int(p[:4]),1,1)+datetime.timedelta(days=int(p[4:])+9)
def inter(draw,dans): return len(set(str(draw))&set(str(x) for x in dans))

BASE = os.path.dirname(os.path.abspath(__file__))
fc3d = json.load(open(os.path.join(BASE, "..", "data", "fc3d-prediction.json"), encoding="utf-8"))
rows = []
for k,v in fc3d.items():
    if not v.get('drawNum'): continue
    if v.get('year'): date = datetime.date(v['year'],v['month'],v['day'])
    else: date = period_to_date(k)
    rows.append({'period':str(k),'date':date,'dans':dans_for(date.year,date.month,date.day),
                 'draw':str(v['drawNum']),'cold':v.get('coldDans')})
rows.sort(key=lambda r:r['period']); rows = rows[-300:]

# ===== 样式 =====
HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
WIN_FILL = PatternFill("solid", fgColor="DCFCE7")    # 中出绿
LOSE_FILL = PatternFill("solid", fgColor="FEE2E2")   # 未中红
SKIP_FILL = PatternFill("solid", fgColor="F3F4F6")   # 不投灰
RESON_FILL = PatternFill("solid", fgColor="FEF3C7")  # 共振黄
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()
ws = wb.active
ws.title = "投注明细"
headers = ["序号","期号","日期","星期","时干3胆","冷号3胆","是否共振","是否投注","投注号码","开奖号码","命中个数","是否中出","单期盈亏(元)","累计盈亏(元)"]
ws.append(headers)
for c in range(1, len(headers)+1):
    cell = ws.cell(row=1, column=c)
    cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = CENTER; cell.border = BORDER

total = 0; bet_cnt = 0; win_cnt = 0
for i, r in enumerate(rows, 1):
    res_flag = bool(set(str(x) for x in r['dans']) & set(str(x) for x in r['cold']))  # 共振
    bet = res_flag  # 共振期投注（投冷号）
    n = inter(r['draw'], r['cold'])
    hit = n in (1,2)
    if bet:
        bet_cnt += 1
        if hit: win_cnt += 1
        profit = 3.5 if hit else -6.3
        total += profit
    else:
        profit = 0
    date_str = r['date'].strftime("%Y-%m-%d")
    wd = WEEK_CN[r['date'].isoweekday()]
    ws.append([i, r['period'], date_str, wd, " ".join(map(str,r['dans'])), " ".join(map(str,r['cold'])),
               "✓" if res_flag else "✗", "投" if bet else "—",
               " ".join(map(str,r['cold'])) if bet else "—",
               r['draw'], n, "✅中出" if hit else ("❌未中" if bet else "—"),
               round(profit,2), round(total,2)])
    ri = ws.max_row
    for c in range(1, len(headers)+1):
        ws.cell(row=ri, column=c).border = BORDER
        ws.cell(row=ri, column=c).alignment = CENTER
    if res_flag:
        for c in range(1, len(headers)+1):
            ws.cell(row=ri, column=c).fill = RESON_FILL
    if bet and hit:
        for c in range(1, len(headers)+1):
            ws.cell(row=ri, column=c).fill = WIN_FILL
    elif bet and not hit:
        for c in range(1, len(headers)+1):
            ws.cell(row=ri, column=c).fill = LOSE_FILL

# 列宽
widths = [6, 10, 12, 7, 12, 12, 9, 9, 12, 10, 9, 10, 13, 13]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# ===== Sheet2 汇总 =====
ws2 = wb.create_sheet("汇总")
ws2.append(["📊 共振期投冷号策略 · 福彩3D 最近300期回测"])
ws2.append([])
rows_sum = [
    ["投注规则", "只在 时干3胆 ∩ 冷号3胆 有交集（共振）的期数投注，投冷号3胆"],
    ["中出规则", "开奖号去重数字与投注号码交集 = 1 或 2 个 → 中出；0 或 3 个（全中）→ 未中"],
    ["投注金额", "每期固定 1 倍 6.3 元，中出得 9.8 元（净 +3.5），未中 -6.3 元"],
    ["", ""],
    ["总期数", len(rows)],
    ["投注期数（共振期）", bet_cnt],
    ["跳过期数（非共振期）", len(rows)-bet_cnt],
    ["中出期数", win_cnt],
    ["未中期数", bet_cnt - win_cnt],
    ["中出率", f"{win_cnt/bet_cnt*100:.1f}%"],
    ["总投入(元)", round(bet_cnt*6.3, 1)],
    ["总回报(元)", round(win_cnt*9.8, 1)],
    ["净盈亏(元)", round(total, 1)],
    ["", ""],
    ["对比：全期投时干", "300期 投入1890元 净+60.2元"],
    ["对比：全期投冷号", "300期 投入1890元 净+177.8元"],
    ["对比：非共振期投时干", "99期 投入623.7元 净-35.7元（越投越亏）"],
]
for r in rows_sum:
    ws2.append(r)
ws2.column_dimensions['A'].width = 28
ws2.column_dimensions['B'].width = 60
for c in range(1, 3):
    ws2.cell(row=1, column=c).font = Font(bold=True, size=14)

out = os.path.join(BASE, "..", "3D共振投冷号300期投注明细.xlsx")
wb.save(out)
print("✅ 已生成:", out)
print(f"总期数 {len(rows)} | 投注 {bet_cnt} 期 | 中出 {win_cnt} | 净盈亏 {round(total,1)} 元")
