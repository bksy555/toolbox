#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 排列三全期投冷号 300期 投注明细 Excel（交集=1才中 / 8.84投19.6中）"""
import datetime, json, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

def period_to_date(p):
    return datetime.date(int(p[:4]), 1, 1) + datetime.timedelta(days=int(p[4:]) + 9)

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
data = json.load(open(os.path.join(BASE, 'data', 'p3-prediction.json'), encoding='utf-8'))
rows = []
for k, v in data.items():
    if not v.get('drawNum') or not v.get('coldDans'):
        continue
    if v.get('year'):
        date = datetime.date(v['year'], v['month'], v['day'])
    else:
        date = period_to_date(k)
    rows.append({'period': str(k), 'date': date, 'draw': str(v['drawNum']), 'cold': v['coldDans']})
rows.sort(key=lambda r: r['period'])
rows = rows[-300:]

COST, PRIZE = 8.84, 19.6
WEEK_CN = ["", "周一", "周二", "周三", "周四", "周五", "周六", "周日"]

def inter(draw, dans):
    return len(set(str(draw)) & set(str(x) for x in dans))

HF = PatternFill("solid", fgColor="4F46E5")
HFONT = Font(color="FFFFFF", bold=True, size=11)
WIN = PatternFill("solid", fgColor="DCFCE7")
LOSE = PatternFill("solid", fgColor="FEE2E2")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Border(*[Side(style="thin", color="D1D5DB")] * 4)

wb = Workbook()
ws = wb.active
ws.title = "策略汇总"
ws.append(["📊 排列三 全期投冷号 回测最近300期"])
ws.append(["规则：预测3码 ∩ 开奖号去重 = 1 个才算中；0/2/3 个都不算中奖"])
ws.append(["赔率：每期投入 8.84 元，中得 19.6 元（净 +10.76）| 盈亏平衡点 45.1%"])
ws.append([])
head = ["玩法", "策略", "投注期数", "中出", "未中", "中出率", "总投入(元)", "总回报(元)", "净盈亏(元)", "每元回报率", "最大连错", "盈亏判定"]
ws.append(head)
for c in range(1, len(head) + 1):
    cell = ws.cell(row=5, column=c); cell.fill = HF; cell.font = HFONT; cell.alignment = CENTER; cell.border = THIN

n = len(rows); wins = 0; ti = 0; to = 0; maxl = cur = 0
for r in rows:
    ti += COST
    if inter(r['draw'], r['cold']) == 1:
        wins += 1; to += PRIZE; cur = 0
    else:
        cur += 1; maxl = max(maxl, cur)
profit = round(to - ti, 2)
ws.append(["排列三", "全期投冷号 ★", n, wins, n - wins, f"{wins/n*100:.1f}%", round(ti, 2), round(to, 2), profit, f"{profit/ti*100:+.1f}%", maxl, "✅ 盈利" if profit > 0 else "❌ 亏损"])
ri = ws.max_row
for c in range(1, len(head) + 1):
    cell = ws.cell(row=ri, column=c); cell.border = THIN; cell.alignment = CENTER; cell.font = Font(bold=True, color="7C3AED")
ws.append([])
ws.append(["说明：该策略为每期都投冷号3胆，不挑期。对比：P3全期投时干 55.3%/+601.6 元（最优）、共振期投冷号 52.1%/+290.8 元"])
for col, w in zip(range(1, len(head) + 1), [10, 14, 12, 8, 8, 10, 13, 13, 13, 13, 10, 10]):
    ws.column_dimensions[chr(64 + col)].width = w

ws2 = wb.create_sheet("P3明细-全期投冷号")
ws2.append(["排列三 全期投冷号 300期投注明细"])
ws2.append(["规则：交集=1才中 | 每期投 8.84 元，中得 19.6 元（净 +10.76）"])
ws2.append([])
head2 = ["序号", "期号", "日期", "星期", "冷号3胆", "开奖号码", "命中个数", "是否中(交集=1)", "单期盈亏(元)", "累计盈亏(元)"]
ws2.append(head2)
for c in range(1, len(head2) + 1):
    cell = ws2.cell(row=4, column=c); cell.fill = HF; cell.font = HFONT; cell.alignment = CENTER; cell.border = THIN
total = 0
for i, r in enumerate(rows, 1):
    hitn = inter(r['draw'], r['cold'])
    hit = hitn == 1
    profit = round(PRIZE - COST, 2) if hit else -COST
    total += profit
    ws2.append([i, r['period'], r['date'].strftime("%Y-%m-%d"), WEEK_CN[r['date'].isoweekday()],
                " ".join(map(str, r['cold'])), r['draw'], hitn,
                "✅中" if hit else "❌", round(profit, 2), round(total, 2)])
    ri = ws2.max_row
    for c in range(1, len(head2) + 1):
        cell = ws2.cell(row=ri, column=c); cell.border = THIN; cell.alignment = CENTER
        cell.fill = WIN if hit else LOSE
for col, w in zip(range(1, len(head2) + 1), [6, 10, 12, 7, 12, 11, 10, 13, 13, 13]):
    ws2.column_dimensions[chr(64 + col)].width = w

out = os.path.join(BASE, "P3全期投冷号300期投注明细.xlsx")
wb.save(out)
print("✅ 已生成:", out)
