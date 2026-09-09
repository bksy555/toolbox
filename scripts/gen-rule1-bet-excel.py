#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 新规则（交集=1才中）回测 300 期 Excel：3D + P3 全策略 + 投注明细"""
import datetime, json, io, os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

GAN_TO_DAN = {"甲":[1,4,8],"乙":[3,4,8],"丙":[3,4,9],"丁":[2,4,9],"戊":[3,4,9],"己":[2,4,9],"庚":[2,7,9],"辛":[2,6,7],"壬":[1,6,7],"癸":[1,6,8]}
TIAN_GAN = ["甲","乙","丙","丁","戊","己","庚","辛","壬","癸"]
WEEK_CN = ["", "周一", "周二", "周三", "周四", "周五", "周六", "周日"]

def get_hai_hour_gan(y,m,d):
    diff = (datetime.date(y,m,d)-datetime.date(1900,1,1)).days
    cycle = ((diff+10)%60+60)%60
    return TIAN_GAN[(((cycle%10)%5)*2+11)%10]
def dans_for(y,m,d): return GAN_TO_DAN[get_hai_hour_gan(y,m,d)]
def period_to_date(p):
    return datetime.date(int(p[:4]),1,1)+datetime.timedelta(days=int(p[4:])+9)
def inter(draw,dans): return len(set(str(draw))&set(str(x) for x in dans))

COST, PRIZE = 8.84, 19.6

def build_rows(path, limit=300):
    data = json.load(io.open(path, encoding='utf-8'))
    rows = []
    for k,v in data.items():
        if not v.get('drawNum'): continue
        if v.get('year'): date = datetime.date(v['year'],v['month'],v['day'])
        else: date = period_to_date(k)
        rows.append({'period':str(k),'date':date,'dans':dans_for(date.year,date.month,date.day),
                     'draw':str(v['drawNum']),'cold':v.get('coldDans')})
    rows.sort(key=lambda r:r['period'])
    return rows[-limit:]

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
r3 = build_rows(os.path.join(BASE,'data','fc3d-prediction.json'))
rp = build_rows(os.path.join(BASE,'data','p3-prediction.json'))

# ===== 样式 =====
HF = PatternFill("solid", fgColor="4F46E5"); HFONT = Font(color="FFFFFF", bold=True, size=11)
WIN = PatternFill("solid", fgColor="DCFCE7"); LOSE = PatternFill("solid", fgColor="FEE2E2")
SKIP = PatternFill("solid", fgColor="F3F4F6"); RESON = PatternFill("solid", fgColor="FEF3C7")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Border(*[Side(style="thin", color="D1D5DB")]*4)

wb = Workbook()

# ===== Sheet1 汇总 =====
ws = wb.active; ws.title = "策略汇总"
def summarize(rows, pick, bet_filter=None, label=''):
    n = wins = 0; ti = to = 0; maxl = cur = 0
    for r in rows:
        dans = r[pick] if pick=='cold' else r['dans']
        if not dans: continue
        if bet_filter and not bet_filter(r): continue
        n += 1; ti += COST
        hit = inter(r['draw'], dans) == 1
        if hit: wins += 1; to += PRIZE; cur = 0
        else: cur += 1; maxl = max(maxl, cur)
    return n, wins, ti, to, maxl

ws.append(["📊 新规则回测：预测3码 ∩ 开奖号去重 = 1 个才算中；0/2/3 个都不算中奖"])
ws.append(["赔率：每期投入 8.84 元，中得 19.6 元（净 +10.76）| 盈亏平衡点 45.1%"])
ws.append(["回测区间：最近 300 期（3D 与 P3 各自独立）"])
ws.append([])
head = ["玩法","策略","投注期数","中出","未中","中出率","总投入(元)","总回报(元)","净盈亏(元)","每元回报率","最大连错","盈亏判定"]
ws.append(head)
for c in range(1, len(head)+1):
    cell = ws.cell(row=5, column=c); cell.fill = HF; cell.font = HFONT; cell.alignment = CENTER; cell.border = THIN

def add_row(game, label, rows, pick, bet_filter=None, star=False):
    n, wins, ti, to, maxl = summarize(rows, pick, bet_filter)
    profit = round(to-ti, 2)
    rate = f"{profit/ti*100:.1f}%" if ti > 0 else "—"
    verdict = "✅ 盈利" if profit > 0 else ("❌ 亏损" if profit < 0 else "➖ 持平")
    vals = [game, label, n, wins, n-wins, f"{wins/n*100:.1f}%" if n else "—",
            round(ti,2), round(to,2), profit, rate, maxl, verdict]
    ws.append(vals)
    ri = ws.max_row
    for c in range(1, len(head)+1):
        cell = ws.cell(row=ri, column=c); cell.border = THIN; cell.alignment = CENTER
        if star: cell.font = Font(bold=True, color="7C3AED")

r3_cold = [r for r in r3 if r['cold']]
r3_res = [r for r in r3_cold if set(str(x) for x in r['dans']) & set(str(x) for x in r['cold'])]
r3_nres = [r for r in r3_cold if not (set(str(x) for x in r['dans']) & set(str(x) for x in r['cold']))]
rp_cold = [r for r in rp if r['cold']]
rp_res = [r for r in rp_cold if set(str(x) for x in r['dans']) & set(str(x) for x in r['cold'])]

for game, rows, cold_rows, res_rows in [("福彩3D", r3, r3_cold, r3_res), ("排列三", rp, rp_cold, rp_res)]:
    add_row(game, "① 全期投时干3胆", rows, 'dans')
    add_row(game, "② 全期投冷号3胆", cold_rows, 'cold')
    add_row(game, "③ 共振期投时干", res_rows, 'dans')
    add_row(game, "④ 共振期投冷号 ★", res_rows, 'cold', star=True)
    add_row(game, "⑤ 非共振期投时干（对照）", r3_nres if game=="福彩3D" else [r for r in rp_cold if not (set(str(x) for x in r['dans']) & set(str(x) for x in r['cold']))], 'dans')
    ws.append([])

ws.column_dimensions['A'].width = 10; ws.column_dimensions['B'].width = 24
for c in range(3, 13): ws.column_dimensions[get_column_letter(c)].width = 14
ws.cell(row=1, column=1).font = Font(bold=True, size=14)

# ===== Sheet2 3D明细（策略=共振期投冷号）=====
def make_detail(ws, rows, title, strategy):
    ws.append([title]); ws.append([strategy])
    ws.append([])
    head = ["序号","期号","日期","星期","时干3胆","冷号3胆","是否共振","是否投注","投注号码","开奖号码","命中个数","是否中(交集=1)","单期盈亏(元)","累计盈亏(元)"]
    ws.append(head)
    for c in range(1, len(head)+1):
        cell = ws.cell(row=4, column=c); cell.fill = HF; cell.font = HFONT; cell.alignment = CENTER; cell.border = THIN
    total = 0
    for i, r in enumerate(rows, 1):
        res_flag = set(str(x) for x in r['dans']) & set(str(x) for x in r['cold'])
        if strategy == '共振投冷号':
            bet = bool(res_flag); pick = 'cold'; skip_note = "投" if bet else "跳过"
        else:  # 全期投时干
            bet = True; pick = 'dans'; skip_note = "投"
        dans = r[pick]
        n = inter(r['draw'], dans)
        hit = n == 1
        if bet:
            profit = round(PRIZE - COST, 2) if hit else -COST
            total += profit
        else:
            profit = 0
        ws.append([i, r['period'], r['date'].strftime("%Y-%m-%d"), WEEK_CN[r['date'].isoweekday()],
                   " ".join(map(str,r['dans'])), " ".join(map(str,r['cold'])) if r['cold'] else "—",
                   "✓" if res_flag else "✗", skip_note, " ".join(map(str,dans)) if bet else "—",
                   r['draw'], n, "✅中" if hit else ("❌" if bet else "—"),
                   round(profit,2), round(total,2)])
        ri = ws.max_row
        for c in range(1, len(head)+1):
            cell = ws.cell(row=ri, column=c); cell.border = THIN; cell.alignment = CENTER
        if strategy == '共振投冷号':
            if res_flag:
                for c in range(1, len(head)+1): ws.cell(row=ri, column=c).fill = RESON
            if bet and hit:
                for c in range(1, len(head)+1): ws.cell(row=ri, column=c).fill = WIN
            elif bet and not hit:
                for c in range(1, len(head)+1): ws.cell(row=ri, column=c).fill = LOSE
            elif not bet:
                for c in range(1, len(head)+1): ws.cell(row=ri, column=c).fill = SKIP
        else:
            for c in range(1, len(head)+1):
                ws.cell(row=ri, column=c).fill = WIN if hit else LOSE
    for col, w in zip(range(1, len(head)+1), [6,10,12,7,12,12,9,9,12,10,9,13,13,13]):
        ws.column_dimensions[get_column_letter(col)].width = w

ws2 = wb.create_sheet("3D明细-共振投冷号")
make_detail(ws2, r3, "福彩3D 最近300期 投注明细", "共振投冷号")
ws3 = wb.create_sheet("P3明细-全期投时干")
make_detail(ws3, rp, "排列三 最近300期 投注明细", "全期投时干")

out = os.path.join(BASE, "3D-P3新规则回测300期投注明细.xlsx")
wb.save(out)
print("✅ 已生成:", out)
