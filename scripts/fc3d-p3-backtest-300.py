#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
3D / P3 预测3码 300 期回测 Excel 生成器
- Sheet1 「3D回测」：最近300期（有号码），标准算法重算时干天干3胆
- Sheet2 「P3回测」：最近300期（有号码），同算法
- Sheet3 「分析」：准确率、最大连错期数、盈亏金额（每期投入6.3，中奖返9.8）
"""
import datetime
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ===== 时干天干3胆算法（与 fc3d-prediction-cron.sh 一致） =====
GAN_TO_DAN = {
    "甲": [1, 4, 8], "乙": [3, 4, 8], "丙": [3, 4, 9],
    "丁": [2, 4, 9], "戊": [3, 4, 9], "己": [2, 4, 9],
    "庚": [2, 7, 9], "辛": [2, 6, 7], "壬": [1, 6, 7], "癸": [1, 6, 8],
}
TIAN_GAN = ["甲", "乙", "丙", "丁", "戊", "己", "庚", "辛", "壬", "癸"]


def get_hai_hour_gan(y, m, d):
    diff = (datetime.date(y, m, d) - datetime.date(1900, 1, 1)).days
    cycle_num = ((diff + 10) % 60 + 60) % 60
    day_gan_idx = cycle_num % 10
    gan_idx = ((day_gan_idx % 5) * 2 + 11) % 10  # 亥时
    return TIAN_GAN[gan_idx]


def dans_for(y, m, d):
    return GAN_TO_DAN[get_hai_hour_gan(y, m, d)]


def calc_result(dans, draw_num):
    if not draw_num:
        return None
    draw_arr = [int(c) for c in str(draw_num)]
    hits = sum(1 for d in dans if d in draw_arr)
    return "✅" if hits >= 1 else "❌"


# ===== 读取数据 =====
BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
d3 = json.load(open(os.path.join(BASE, "data", "fc3d-prediction.json"), encoding="utf-8"))
p3 = json.load(open(os.path.join(BASE, "data", "p3-prediction.json"), encoding="utf-8"))

WEEK_CN = ["", "周一", "周二", "周三", "周四", "周五", "周六", "周日"]


def period_to_date(period):
    """期号反推日期：期号 = 年内第几天 - 10"""
    p = str(period)
    year = int(p[:4])
    day_of_year = int(p[4:]) + 10
    return datetime.date(year, 1, 1) + datetime.timedelta(days=day_of_year - 1)


def build_rows(data, limit=300):
    """取最近 limit 期（有号码），标准算法重算 dans"""
    rows = []
    for k, v in data.items():
        if not v.get("drawNum"):
            continue
        if v.get("year"):
            date = datetime.date(v["year"], v["month"], v["day"])
        else:
            date = period_to_date(k)
        rows.append({
            "period": k,
            "date": date,
            "dans": dans_for(date.year, date.month, date.day),
            "draw": str(v["drawNum"]),
        })
    rows.sort(key=lambda r: r["period"])
    rows = rows[-limit:]
    return rows


d3_rows = build_rows(d3, 300)
p3_rows = build_rows(p3, 300)
print(f"3D 回测期数: {len(d3_rows)} ({d3_rows[0]['period']}~{d3_rows[-1]['period']})")
print(f"P3 回测期数: {len(p3_rows)} ({p3_rows[0]['period']}~{p3_rows[-1]['period']})")

# ===== 样式 =====
HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
WIN_FILL = PatternFill("solid", fgColor="DCFCE7")
LOSE_FILL = PatternFill("solid", fgColor="FEE2E2")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

wb = Workbook()


def write_sheet(ws, title, rows):
    headers = ["期号", "日期", "星期", "预测3胆", "中奖号码", "命中", "单期盈亏(元)", "累计盈亏(元)"]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    total = 0
    for i, r in enumerate(rows):
        res = calc_result(r["dans"], r["draw"])
        profit = 9.8 - 6.3 if res == "✅" else -6.3
        total += profit
        date_str = r["date"].strftime("%Y-%m-%d")
        weekday = WEEK_CN[r["date"].isoweekday()]
        dans_str = " ".join(map(str, r["dans"]))
        ws.append([r["period"], date_str, weekday, dans_str, r["draw"], res, round(profit, 2), round(total, 2)])
        row_idx = ws.max_row
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=c)
            cell.alignment = CENTER
            cell.border = BORDER
        if res == "✅":
            ws.cell(row=row_idx, column=6).fill = WIN_FILL
            ws.cell(row=row_idx, column=7).fill = WIN_FILL
        else:
            ws.cell(row=row_idx, column=6).fill = LOSE_FILL
            ws.cell(row=row_idx, column=7).fill = LOSE_FILL
        ws.cell(row=row_idx, column=7).number_format = "+0.00;-0.00;0.00"
        ws.cell(row=row_idx, column=8).number_format = "+0.00;-0.00;0.00"

    ws.freeze_panes = "A2"
    widths = [12, 14, 10, 18, 14, 10, 16, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return total, rows


def compute_stats(rows, price=6.3, prize=9.8):
    """计算准确率/最大连错/盈亏"""
    total_win = 0
    max_consecutive_lose = 0
    cur_lose = 0
    cum = 0.0
    results = []
    for r in rows:
        res = calc_result(r["dans"], r["draw"])
        results.append(res)
        if res == "✅":
            total_win += 1
            cur_lose = 0
            cum += prize - price
        else:
            cur_lose += 1
            max_consecutive_lose = max(max_consecutive_lose, cur_lose)
            cum -= price
    n = len(rows)
    accuracy = total_win / n if n else 0
    total_invest = n * price
    total_prize = total_win * prize
    net = cum
    return {
        "n": n,
        "win": total_win,
        "lose": n - total_win,
        "accuracy": accuracy,
        "max_lose": max_consecutive_lose,
        "invest": total_invest,
        "prize": total_prize,
        "net": net,
    }


# ===== Sheet1: 3D 回测 =====
ws1 = wb.active
ws1.title = "3D回测"
total_d3, _ = write_sheet(ws1, "3D 回测", d3_rows)

# ===== Sheet2: P3 回测 =====
ws2 = wb.create_sheet("P3回测")
total_p3, _ = write_sheet(ws2, "P3 回测", p3_rows)

# ===== Sheet3: 分析 =====
ws3 = wb.create_sheet("分析")
s3 = compute_stats(d3_rows)
s4 = compute_stats(p3_rows)

ws3.merge_cells("A1:H1")
ws3["A1"] = "🎯 预测3码 300 期回测分析（每期投入 1 倍 = 6.3 元，中奖返 9.8 元）"
ws3["A1"].font = Font(bold=True, size=13, color="4F46E5")
ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws3.row_dimensions[1].height = 28

rows_analysis = [
    ["项目", "福彩3D（300期）", "排列三（300期）"],
    ["回测区间", f"{d3_rows[0]['period']} ~ {d3_rows[-1]['period']}", f"{p3_rows[0]['period']} ~ {p3_rows[-1]['period']}"],
    ["回测总期数", s3["n"], s4["n"]],
    ["✅ 命中期数", s3["win"], s4["win"]],
    ["❌ 未中期数", s3["lose"], s4["lose"]],
    ["🎯 准确率", f"{s3['accuracy']*100:.2f}%", f"{s4['accuracy']*100:.2f}%"],
    ["⚠️ 最大连错期数", s3["max_lose"], s4["max_lose"]],
    ["💸 总投入(6.3/期)", f"{s3['invest']:.1f} 元", f"{s4['invest']:.1f} 元"],
    ["💰 总返奖(9.8/中)", f"{s3['prize']:.1f} 元", f"{s4['prize']:.1f} 元"],
    ["📈 净盈亏", f"{s3['net']:+.1f} 元", f"{s4['net']:+.1f} 元"],
]
for r in rows_analysis:
    ws3.append(r)
for row_idx in range(2, 2 + len(rows_analysis)):
    for col_idx in range(1, 4):
        cell = ws3.cell(row=row_idx, column=col_idx)
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws3.cell(row=row_idx, column=1).font = Font(bold=True)
ws3.cell(row=2, column=1).font = Font(bold=True, color="4F46E5")
ws3.cell(row=2, column=2).font = Font(bold=True, color="4F46E5")
ws3.cell(row=2, column=3).font = Font(bold=True, color="4F46E5")

ws3.merge_cells("A13:H13")
ws3["A13"] = "📌 说明：预测3码为当日时干天干算法（3D 与 P3 同算法，同日组合一致）；命中 = 预测3胆中至少1个数字出现在开奖号码中；盈亏按每期固定投入 6.3 元、命中期返 9.8 元计算（净 +3.5 / 错 -6.3），未计税费。"
ws3["A13"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws3["A13"].font = Font(size=10, color="6B7280")
ws3.row_dimensions[13].height = 44
ws3.column_dimensions["A"].width = 24
ws3.column_dimensions["B"].width = 26
ws3.column_dimensions["C"].width = 26

# 输出
out = os.path.join(BASE, "data", "回测报告_3D_P3_300期.xlsx")
wb.save(out)
print(f"\n✅ 已生成: {out}")
print(f"3D 净盈亏: {total_d3:+.1f} 元 | P3 净盈亏: {total_p3:+.1f} 元")
print(f"3D 准确率: {s3['accuracy']*100:.2f}% 最大连错 {s3['max_lose']} 期")
print(f"P3 准确率: {s4['accuracy']*100:.2f}% 最大连错 {s4['max_lose']} 期")